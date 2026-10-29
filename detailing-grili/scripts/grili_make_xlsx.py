# GRILI.UA — генератор Excel. Після генерації формули обовʼязково перераховуються (scripts/build.sh це робить). Дані — з grili_data.py у робочій папці, логотип — assets/grili_logo.png
import os, sys, glob
HERE = os.path.dirname(os.path.abspath(__file__))
# Дані (grili_data.py) беруться з робочої папки (або з GRILI_DATA_DIR), а не зі скілла
sys.path.insert(0, os.environ.get("GRILI_DATA_DIR", os.getcwd()))
LOGO = os.environ.get("GRILI_LOGO", os.path.join(HERE, "..", "assets", "grili_logo.png"))
OUT = os.environ.get("GRILI_OUT", ".")
try:
    from grili_data import DATA, COMPANY, GRAPHITE, ORANGE, GREEN
except ImportError:
    sys.exit("Не знайдено grili_data.py у робочій папці. Скопіюйте assets/grili_data_template.py -> grili_data.py і заповніть.")
PREFIX = DATA.get("file_prefix", "")  # напр. "Зі знижкою — "
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.worksheet.page import PageMargins
from PIL import Image as PILImage

FONT = "Calibri"
def f(size=10, bold=False, color=GRAPHITE):
    return Font(name=FONT, size=size, bold=bold, color=color)
fill = lambda c: PatternFill("solid", start_color=c)
thin = Side(style="thin", color="C9CFD1")
med = Side(style="medium", color=GRAPHITE)
orange_side = Side(style="thick", color=ORANGE)

wb = Workbook()
ws = wb.active
ws.title = "Пропозиція"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [5, 62, 14, 11, 18]):
    ws.column_dimensions[col].width = w

# --- шапка: логотип + контакти ---
pil = PILImage.open(LOGO)
pil.thumbnail((700, 700))
pil.save("logo_xl.png")
img = XLImage("logo_xl.png")
img.width = 250
img.height = int(250 * pil.height / pil.width)
ws.add_image(img, "A1")
for r in range(1, 5):
    ws.row_dimensions[r].height = 15
ws.merge_cells("C1:E4")
ws["C1"] = ("   ".join(COMPANY["phones"][:1]) + "\n" + "   ".join(COMPANY["phones"][1:]) + "\n"
            + COMPANY["email"] + "\n" + "; ".join(COMPANY["shops"]))
ws["C1"].font = f(8.5)
ws["C1"].alignment = Alignment(horizontal="right", vertical="top", wrap_text=True)
for c in "ABCDE":
    ws[f"{c}5"].border = Border(bottom=orange_side)
ws.row_dimensions[5].height = 6

# --- блок замовника ---
ws["A7"] = f"Пропозиція №{DATA['number']}"
ws["A7"].font = f(15, True)
ws["A8"] = DATA["title"]
ws["A8"].font = f(11)
ws["A9"] = f"від {DATA['date']}"
ws["A9"].font = f(9, color=ORANGE)
for r, (lab, val) in zip((7, 8, 9), [(DATA.get("place_label","Місце монтажу"), DATA["place"]), ("Замовник", DATA["customer"]),
                                      ("Конт. тел.", DATA["phone"] or "__________")]):
    ws.merge_cells(f"C{r}:E{r}")
    ws[f"C{r}"] = f"{lab}: {val}"
    ws[f"C{r}"].font = f(10, lab != "")
    ws[f"C{r}"].alignment = Alignment(horizontal="right")
ws.row_dimensions[7].height = 22

# --- таблиця ---
HR = 11
for c, t in zip("ABCDE", ["№", "Найменування", "Вартість, грн", "Кількість", "Кінцевий прорахунок, грн"]):
    cell = ws[f"{c}{HR}"]
    cell.value = t
    cell.font = f(10, True, "FFFFFF")
    cell.fill = fill(GRAPHITE)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
ws.row_dimensions[HR].height = 30

row = HR + 1
n = 0
sub_cells = []
for sec, items in DATA["sections"]:
    if not items:
        continue
    ws.merge_cells(f"A{row}:E{row}")
    ws[f"A{row}"] = sec
    ws[f"A{row}"].font = f(10, True, "FFFFFF")
    for c in "ABCDE":
        ws[f"{c}{row}"].fill = fill(GREEN)
    ws.row_dimensions[row].height = 18
    row += 1
    first = row
    priced = False
    for it in items:
        name, price, qty = it[0], it[1], it[2]
        extra = it[3] if len(it) > 3 else None
        n += 1
        ws[f"A{row}"] = n
        ws[f"B{row}"] = name
        if price is None:
            # рядок без суми: текст замість вартості, у підсумки не входить
            ws.merge_cells(f"C{row}:E{row}")
            ws[f"C{row}"] = extra or ""
        else:
            priced = True
            ws[f"C{row}"] = price
            ws[f"D{row}"] = qty
            ws[f"E{row}"] = f"=ROUND(C{row}*D{row},0)"
        for c in "ABCDE":
            cell = ws[f"{c}{row}"]
            cell.font = f(10)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="center", wrap_text=(c in "BC"),
                                       horizontal={"A": "center", "D": "center"}.get(c, "right" if c in "CE" else "left"))
        if price is None:
            ws[f"C{row}"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws[f"C{row}"].number_format = "#,##0"
        ws[f"E{row}"].number_format = "#,##0"
        ws[f"D{row}"].number_format = "0.##"
        lines = max(1, -(-len(name) // 70), 2 if price is None else 1)
        ws.row_dimensions[row].height = 15 * lines + 4
        row += 1
        if extra and price is not None:
            # повна вартість (без знижки) — окремим рядком під позицією, закреслено
            ws[f"B{row}"] = "повна вартість, грн"
            ws[f"B{row}"].font = Font(name=FONT, size=8.5, italic=True, color=GRAPHITE)
            ws[f"B{row}"].alignment = Alignment(horizontal="right", vertical="top")
            ws[f"C{row}"] = extra
            ws[f"C{row}"].font = Font(name=FONT, size=8.5, strike=True, color=GRAPHITE)
            ws[f"C{row}"].number_format = "#,##0"
            ws[f"C{row}"].alignment = Alignment(horizontal="right", vertical="top")
            for c in "ABCDE":
                ws[f"{c}{row}"].border = Border(bottom=thin)
            # верхній рядок позиції без нижньої лінії, щоб пара читалась як одна позиція
            for c in "ABCDE":
                ws[f"{c}{row-1}"].border = Border()
            ws.row_dimensions[row].height = 13
            row += 1
    last = row - 1
    if priced:
        ws.merge_cells(f"A{row}:D{row}")
        ws[f"A{row}"] = f"Разом за розділ «{sec}»:"
        ws[f"A{row}"].font = f(10, True)
        ws[f"A{row}"].alignment = Alignment(horizontal="right")
        # у підсумок входять лише клітинки E позицій (рядки «повна вартість» мають порожню E)
        ws[f"E{row}"] = f"=SUM(E{first}:E{last})"
        ws[f"E{row}"].font = f(10, True, ORANGE)
        ws[f"E{row}"].number_format = "#,##0"
        for c in "ABCDE":
            ws[f"{c}{row}"].border = Border(top=med, bottom=med)
        sub_cells.append(f"E{row}")
        row += 2
    else:
        row += 1

# --- підсумки ---
tot_row = row
ws.merge_cells(f"A{row}:D{row}")
ws[f"A{row}"] = "Разом, грн:"
ws[f"A{row}"].font = f(12, True)
ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
ws[f"E{row}"] = "=" + "+".join(sub_cells)
ws[f"E{row}"].font = f(12, True, "FFFFFF")
ws[f"E{row}"].fill = fill(ORANGE)
ws[f"E{row}"].number_format = "#,##0"
ws[f"E{row}"].alignment = Alignment(horizontal="right", vertical="center")
ws.row_dimensions[row].height = 22
row += 1
ws.merge_cells(f"A{row}:C{row}")
ws[f"A{row}"] = f"Разом, USD (курс НБУ на {DATA['rate_date']}):"
ws[f"A{row}"].font = f(10.5)
ws[f"A{row}"].alignment = Alignment(horizontal="right", vertical="center")
ws[f"D{row}"] = DATA["rate"]
ws[f"D{row}"].number_format = "0.00"
ws[f"D{row}"].font = f(10.5)
ws[f"D{row}"].alignment = Alignment(horizontal="center", vertical="center")
ws[f"E{row}"] = f"=ROUND(E{tot_row}/D{row},0)"
ws[f"E{row}"].font = f(11, True)
ws[f"E{row}"].number_format = "#,##0"
ws[f"E{row}"].border = Border(bottom=Side(style="medium", color=ORANGE))
ws.row_dimensions[row].height = 20
row += 2

# --- умови ---
ws[f"A{row}"] = "Умови"
ws[f"A{row}"].font = f(11, True, GREEN)
row += 1
for ctext in DATA["conditions"]:
    t = ctext.format(**DATA)
    ws[f"A{row}"] = "•"
    ws[f"A{row}"].font = f(10, True, ORANGE)
    ws[f"A{row}"].alignment = Alignment(horizontal="center", vertical="top")
    ws.merge_cells(f"B{row}:E{row}")
    ws[f"B{row}"] = t
    ws[f"B{row}"].font = f(9.5)
    ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[row].height = 14 * max(1, -(-len(t) // 105)) + 2
    row += 1
row += 2

# --- описи товарів (необов'язково) ---
for dtitle, drows in DATA.get("descriptions", []):
    ws[f"A{row}"] = dtitle
    ws[f"A{row}"].font = f(11, True, GREEN)
    row += 1
    for i, (a, b) in enumerate(drows):
        ws.merge_cells(f"A{row}:B{row}")
        ws.merge_cells(f"C{row}:E{row}")
        ws[f"A{row}"] = a
        ws[f"A{row}"].font = f(9.5, True)
        ws[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"C{row}"] = b
        ws[f"C{row}"].font = f(9.5)
        ws[f"C{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        if i % 2 == 1:
            for c in "ABCDE":
                ws[f"{c}{row}"].fill = fill("EEF1F1")
        ws.row_dimensions[row].height = 14 * max(1, -(-len(b) // 60), -(-len(a) // 45)) + 3
        row += 1
    row += 2

# --- креслення (необов'язково) ---
from openpyxl.worksheet.pagebreak import Break
for gtitle, gpath in DATA.get("drawings", []):
    ws.row_breaks.append(Break(id=row - 1))  # креслення — з нової сторінки, щоб не розрізалось
    ws[f"A{row}"] = gtitle
    ws[f"A{row}"].font = f(11, True, GREEN)
    row += 1
    gp = PILImage.open(gpath)
    total_w_px = sum((ws.column_dimensions[c].width or 8.43) for c in "ABCDE") * 7
    gw = int(total_w_px)
    gh = int(gw * gp.height / gp.width)
    gimg = XLImage(gpath)
    gimg.width, gimg.height = gw, gh
    ws.add_image(gimg, f"A{row}")
    nrows = -(-gh // 20)
    for r in range(row, row + nrows):
        ws.row_dimensions[r].height = 15
    row += nrows + 2

# --- підпис ---
for lab1, lab2 in [("Менеджер:", "Підпис:"), ("Телефон:", "Дата:")]:
    ws[f"A{row}"] = lab1 + " " + "_" * 36
    ws[f"A{row}"].font = f(10.5, True)
    ws[f"C{row}"] = lab2 + " " + "_" * 22
    ws[f"C{row}"].font = f(10.5, True)
    ws.row_dimensions[row].height = 22
    row += 1

# --- друк ---
ws.print_area = f"A1:E{row}"
ws.print_title_rows = f"{HR}:{HR}"
ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 0
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_margins = PageMargins(left=0.5, right=0.5, top=0.7, bottom=0.6, header=0.3, footer=0.3)
ws.HeaderFooter.differentFirst = True
ws.oddHeader.left.text = "GRILI.UA — все для барбекю"
ws.oddHeader.left.color = GREEN
ws.oddHeader.right.text = "0 800 330 287  •  (067) 714 28 82  •  (050) 495 34 06  •  info@grili.com.ua"
ws.oddHeader.right.size = 8
ws.oddFooter.left.text = f"Пропозиція №{DATA['number']}"
ws.oddFooter.left.size = 8
ws.oddFooter.right.text = "Сторінка &P з &N"
ws.oddFooter.right.size = 8
ws.firstFooter.left.text = f"Пропозиція №{DATA['number']}"
ws.firstFooter.left.size = 8
ws.firstFooter.right.text = "Сторінка &P з &N"
ws.firstFooter.right.size = 8

wb.active = 0
wb.save(f"{OUT}/{PREFIX}Пропозиція №{DATA['number']} — {DATA['work_type']} — {DATA['customer']}.xlsx")
print("ok")
