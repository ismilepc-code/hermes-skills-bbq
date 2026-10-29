#!/usr/bin/env python3
"""Перевірка перерахованого Excel: немає помилок, суми збігаються з grili_data.py.
Використання (з робочої папки з grili_data.py): python3 verify_xlsx.py <файл.xlsx>"""
import os, sys
sys.path.insert(0, os.environ.get("GRILI_DATA_DIR", os.getcwd()))
from openpyxl import load_workbook
from grili_data import DATA

ws = load_workbook(sys.argv[1], data_only=True).active
errs = [c.coordinate for r in ws.iter_rows() for c in r if isinstance(c.value, str) and c.value.startswith("#")]
exp = sum(round(it[1] * it[2]) for _, items in DATA["sections"] for it in items if it[1] is not None)
exp_usd = round(exp / DATA["rate"]) if DATA["rate"] else None
got = got_usd = None
for r in ws.iter_rows():
    a = r[0].value
    if a == "Разом, грн:":
        got = r[4].value
    elif isinstance(a, str) and a.startswith("Разом, USD"):
        got_usd = r[4].value
print(f"помилки у клітинках: {errs or 'немає'}")
print(f"Разом, грн:  очікувалось {exp}, в Excel {got}")
print(f"Разом, USD:  очікувалось {exp_usd}, в Excel {got_usd}")
ok = not errs and got == exp and got_usd == exp_usd
print("OK" if ok else "РОЗБІЖНІСТЬ — перевірте файл!")
sys.exit(0 if ok else 1)
