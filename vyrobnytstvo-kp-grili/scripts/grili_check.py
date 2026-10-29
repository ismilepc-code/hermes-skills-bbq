# ШАБЛОН GRILI.UA — перевірка перед видачею.
# Запускати в робочій папці ПІСЛЯ генерації та перерахунку Excel (recalc.py).
# Перевіряє: номер = дата ДДММРР, термін дії = дата + 3 робочі дні, курс на дату подачі (2 знаки),
# суми розділів / «Разом, грн» / «Разом, USD» у Excel збігаються з тим, що рахує PDF, і немає помилок формул.
import os
import sys
import datetime as dt
from grili_data import DATA

OUT = os.environ.get("GRILI_OUT", ".")
problems = []

def d(s):
    return dt.datetime.strptime(s, "%d.%m.%Y").date()

def plus_workdays(day, n):
    while n:
        day += dt.timedelta(days=1)
        if day.weekday() < 5:  # пн–пт
            n -= 1
    return day

# --- номер, дати, курс ---
date = d(DATA["date"])
num = DATA["number"].split("_")[0]
if num != date.strftime("%d%m%y"):
    problems.append(f"Номер {DATA['number']} не відповідає даті {DATA['date']} (очікується {date.strftime('%d%m%y')})")
exp = plus_workdays(date, 3)
if d(DATA["valid_until"]) != exp:
    problems.append(f"Термін дії {DATA['valid_until']} — має бути {exp.strftime('%d.%m.%Y')} (дата + 3 робочі дні)")
if DATA["rate_date"] != DATA["date"]:
    problems.append(f"Курс НБУ взято на {DATA['rate_date']}, а дата подачі {DATA['date']}")
if round(DATA["rate"], 2) != DATA["rate"]:
    problems.append(f"Курс {DATA['rate']} має бути округлений до двох знаків після коми")

# --- суми так само, як у PDF ---
subs = {}
for sec, items in DATA["sections"]:
    s, priced = 0, False
    for it in items:
        if it[1] is not None:
            priced = True
            s += round(it[1] * it[2])
    if items and priced:
        subs[sec] = s
grand = sum(subs.values())
usd = round(grand / DATA["rate"])

# --- файли креслень ---
for _, gpath in DATA.get("drawings", []):
    if not os.path.exists(gpath):
        problems.append(f"Не знайдено файл креслення: {gpath}")

# --- Excel ---
xlsx = f"{OUT}/Пропозиція №{DATA['number']} — {DATA['work_type']} — {DATA['customer']}.xlsx"
if not os.path.exists(xlsx):
    problems.append(f"Не знайдено Excel: {xlsx}")
else:
    from openpyxl import load_workbook
    ws = load_workbook(xlsx, data_only=True)["Пропозиція"]
    x_subs, x_grand, x_usd = {}, None, None
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c.value, str) and c.value.startswith("#"):
                problems.append(f"Помилка формули в {c.coordinate}: {c.value}")
        a = ws[f"A{row[0].row}"].value
        e = ws[f"E{row[0].row}"].value
        if isinstance(a, str):
            if a.startswith("Разом за розділ «"):
                x_subs[a[len("Разом за розділ «"):].rstrip("»:")] = e
            elif a == "Разом, грн:":
                x_grand = e
            elif a.startswith("Разом, USD"):
                x_usd = e
    if x_grand is None:
        problems.append("У Excel немає значення «Разом, грн» — чи запущено recalc.py?")
    for sec, v in subs.items():
        if x_subs.get(sec) != v:
            problems.append(f"Розділ «{sec}»: PDF {v}, Excel {x_subs.get(sec)}")
    if x_grand is not None and x_grand != grand:
        problems.append(f"Разом, грн: PDF {grand}, Excel {x_grand}")
    if x_usd is not None and x_usd != usd:
        problems.append(f"Разом, USD: PDF {usd}, Excel {x_usd}")

fmt = lambda n: f"{n:,}".replace(",", " ")
print("Розділи:", "; ".join(f"{k} — {fmt(v)}" for k, v in subs.items()))
print(f"Разом, грн: {fmt(grand)}   Разом, USD: {fmt(usd)} (курс {DATA['rate']})")
print(f"Дійсна до: {DATA['valid_until']}")
if problems:
    print("\nЗНАЙДЕНО ПРОБЛЕМИ:")
    for p in problems:
        print(" -", p)
    sys.exit(1)
print("\nOK: PDF і Excel збігаються, дати й номер правильні.")
